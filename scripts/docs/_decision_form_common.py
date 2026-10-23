# -*- coding: utf-8 -*-
"""待决表 Word 文档的公共构件（生成 / 回读 / 验收三边共用）。

从 scripts/docs/build_pending_decision_form.py 与
scripts/docs/read_pending_decision_form.py 中抽出，避免两份表各写一遍样式与
勾选框解析逻辑 —— 上一轮踩过的坑（同一行里多个并列勾选框被当成「整行勾上」）
就出在解析逻辑各写各的。

依赖：python-docx（装在 D:\\dev\\python\\python.exe，不写 C 盘）。
"""
import os
import re
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT_BODY = "等线"
FONT_HEAD = "微软雅黑"
C_DARK = RGBColor(0x1F, 0x38, 0x64)
C_MID = RGBColor(0x2E, 0x53, 0x95)
C_GRAY = RGBColor(0x59, 0x59, 0x59)
C_HDRBG = "DCE6F1"
C_LINE = "BFCEDD"

BOX = "□"          # 未勾选
CHECK_MARKS = ["☑", "√", "✓", "✔", "■", "▣", "☒", "x", "X"]
BOX_CHARS = "□☑√✓✔■▣☒"
BOX_RE = re.compile("[" + BOX_CHARS + "]")


# ----------------------------------------------------------------- 排版工具
def set_run_font(run, name=FONT_BODY, size=10.5, bold=False, color=None, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:eastAsia"), name)
    rf.set(qn("w:ascii"), name)
    rf.set(qn("w:hAnsi"), name)


def para(doc, text="", size=10.5, bold=False, color=None, align=None,
         space_before=0, space_after=4, indent=0, name=FONT_BODY, line=1.4, style=None):
    p = doc.add_paragraph(style=style)
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line
    if indent:
        pf.left_indent = Cm(indent)
    if align is not None:
        p.alignment = align
    if text:
        set_run_font(p.add_run(text), name=name, size=size, bold=bold, color=color)
    return p


def heading(doc, text, level=1):
    cfg = {1: (16, 14, 6, C_DARK, FONT_HEAD),
           2: (13, 12, 5, C_MID, FONT_HEAD),
           3: (11.5, 10, 4, RGBColor(0x33, 0x33, 0x33), FONT_HEAD)}
    size, sb, sa, color, font = cfg[level]
    return para(doc, text, size=size, bold=True, color=color,
                space_before=sb, space_after=sa, name=font, line=1.2)


def shade(cell, hexcolor):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcpr.append(shd)


def cell_text(cell, text, size=9, bold=False, color=None, align=None, name=FONT_BODY):
    cell.text = ""
    p = cell.paragraphs[0]
    pf = p.paragraph_format
    pf.space_before = Pt(1.5)
    pf.space_after = Pt(1.5)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15
    if align is not None:
        p.alignment = align
    set_run_font(p.add_run(text), name=name, size=size, bold=bold, color=color)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    trpr.append(h)


def no_split(row):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


def table_borders(table, color=C_LINE, sz=4):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:%s" % edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        borders.append(e)
    tbl_pr.append(borders)


def set_widths(table, widths):
    table.autofit = False
    table.allow_autofit = False
    for row in table.rows:
        for idx, w in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = Cm(w)


def add_page_footer(section, left_text):
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    set_run_font(p.add_run(left_text + "　·　第 "), size=8, color=C_GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    r2 = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "16")
    rpr.append(sz)
    r2.append(rpr)
    t2 = OxmlElement("w:t")
    t2.text = "1"
    r2.append(t2)
    fld.append(r2)
    p._p.append(fld)
    set_run_font(p.add_run(" 页"), size=8, color=C_GRAY)


def blank_lines(doc, n=3, label=None):
    if label:
        para(doc, label, size=9.5, bold=True, color=C_MID, space_before=6, space_after=3)
    for _ in range(n):
        p = para(doc, "", size=10.5, space_after=6)
        ppr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "C9C9C9")
        pbdr.append(bottom)
        ppr.append(pbdr)


def new_document(footer_text):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT_BODY
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)
    add_page_footer(sec, footer_text)
    return doc


# ----------------------------------------------------------------- 勾选框解析
def is_checked(text):
    """这一行是不是被勾了（只看行首若干字符，避免正文里的 √ 误判）。"""
    if not text:
        return False
    for m in CHECK_MARKS:
        if m in text[:8]:
            return True
    if re.search(r"\[[xX√✓]\s*\]", text[:8]):
        return True
    return False


def strip_box(text):
    t = text.strip()
    t = re.sub(r"^[\s　]*[□☑√✓✔■▣☒]?[\s　]*", "", t)
    if is_checked(text):
        t = re.sub(r"^[\[\(]?\s*[xX√✓]?\s*[\]\)]?\s*", "", t)
    return t.strip().strip(".").strip()


def split_boxes(text):
    """把一行拆成若干个「框 + 选项」。

    一行里可能有并列的多个框（例如「□ 推进　□ 暂缓」）。
    必须**逐框**判定勾选状态 —— 只要行内有一个框被勾就整行算勾上，是错的。
    返回 [(行前缀, 选项文本, 是否勾上)]，行内第一个框带上行前缀作为分组名。
    """
    hits = list(BOX_RE.finditer(text))
    if not hits:
        return []
    prefix = text[:hits[0].start()].strip()
    out = []
    for i, h in enumerate(hits):
        end = hits[i + 1].start() if i + 1 < len(hits) else len(text)
        label = text[h.end():end].strip().lstrip(".、 ").strip()
        out.append((prefix, label, h.group(0) in CHECK_MARKS))
    return out


# ----------------------------------------------------------------- 结构验收
def verify(path):
    """结构验收：zip 完整性 + XML 部件可解析 + 关键结构在位。"""
    out = {"file": path, "exists": os.path.exists(path)}
    if not out["exists"]:
        return out
    out["bytes"] = os.path.getsize(path)
    with zipfile.ZipFile(path) as zf:
        bad = zf.testzip()
        out["zip_ok"] = bad is None
        names = zf.namelist()
        out["xml_parts"] = len([n for n in names if n.endswith(".xml")])
        import lxml.etree as ET
        fail = []
        for n in names:
            if n.endswith(".xml"):
                try:
                    ET.fromstring(zf.read(n))
                except Exception as e:
                    fail.append("%s: %s" % (n, e))
        out["xml_parse_failures"] = fail
    doc = Document(path)
    out["paragraphs"] = len(doc.paragraphs)
    out["tables"] = len(doc.tables)
    out["table_rows"] = [len(t.rows) for t in doc.tables]
    out["table_cells"] = [sum(len(r.cells) for r in t.rows) for t in doc.tables]
    # 页脚 PAGE 域
    #
    # PAGE 域有两种等价形态，只认一种会误判：
    #   ① 简单域 <w:fldSimple w:instr=" PAGE ">  —— python-docx 生成
    #   ② 复杂域 fldChar begin / instrText " PAGE " / separate / end —— Word / WPS 重存后
    # 马哥用 WPS 存过的表就是 ②，原判定会把好文档报成「结构验收有问题」。
    # 实测页脚可见文本里页码占位仍在（…第 1 页），所以按「域在位 + 页脚含页码占位」判。
    xml = doc.sections[0].footer._element.xml
    simple = bool(re.search(r'w:instr="\s*PAGE\s*"', xml))
    complex_ = bool(re.search(r"<w:instrText[^>]*>\s*PAGE\s*</w:instrText>", xml))
    footer_text = " ".join(p.text for p in doc.sections[0].footer.paragraphs)
    out["footer_page_field"] = (simple or complex_) and "页" in footer_text
    out["footer_page_field_form"] = "fldSimple" if simple else ("complex" if complex_ else "missing")
    return out


def structure_ok(info):
    return bool(info.get("zip_ok")
                and not info.get("xml_parse_failures")
                and info.get("footer_page_field"))


# ----------------------------------------------------------------- 回读解析
def parse_sections(path, sec_re, ids, global_heading, end_heading, table_rows=None):
    """把一份待决表解析成结构化结果。

    sec_re        : 小节标题正则（如 r"^(D[1-5])[\\s　]"）
    ids           : 小节编号集合，如 {"D1", ..., "D5"}
    global_heading: 总批注区标题前缀
    end_heading   : 结尾区标题前缀
    table_rows    : 速览表应有的行数（含表头）；给了就顺手把「我的选择」列也读出来
    """
    doc = Document(path)
    sections = {}
    cur = None
    state = None
    global_notes = []
    in_global = False
    global_first = True

    # end_heading 可以是单个标题，也可以是若干个标题（元组）。
    # 为什么要兼容多个：待决表从「空白待决」转成「已确认记录」后，
    # 结尾标题从「4　填完之后我会做什么」改成了「4　确认落地清单」，
    # 而仓库里还留着当初的空白表 —— 两种表都要能读回，
    # 只认一个标题会让另一种表的正文漏进总批注区。
    if isinstance(end_heading, str):
        end_headings = (end_heading,)
    else:
        end_headings = tuple(end_heading)

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        m = sec_re.match(text)
        if m and m.group(1) in ids:
            cur = m.group(1)
            sections.setdefault(cur, {"context": "", "recommend": "", "options": [],
                                      "checked": [], "unchecked": [], "choice": "", "notes": []})
            state = "ctx"
            in_global = False
            continue
        if text.startswith(global_heading):
            in_global = True
            global_first = True
            cur = None
            continue
        if any(text.startswith(h) for h in end_headings):
            in_global = False
            continue

        if in_global:
            # 标题后的第一段是填写说明，其余才是他写的总规则
            if global_first:
                global_first = False
                continue
            global_notes.append(text)
            continue

        if cur is None:
            continue

        if text.startswith("小节批注"):
            continue
        if text.startswith("→"):
            sections[cur]["choice"] = text.split("：", 1)[-1].strip().strip("_").strip()
            state = "notes"
            continue
        if text.startswith("💡"):
            sections[cur]["recommend"] = text[1:].strip()
            state = "opt"
            continue
        if BOX in text or is_checked(text):
            for prefix, label, checked in split_boxes(text):
                if not label:
                    continue
                sections[cur]["options"].append(
                    {"group": prefix, "label": label, "checked": checked})
                key = "checked" if checked else "unchecked"
                sections[cur][key].append(label)
            state = "opt"
            continue
        if state == "ctx":
            sections[cur]["context"] = text
            state = "opt"
            continue
        if state == "notes":
            sections[cur]["notes"].append(text)

    # 速览表里的「我的选择」列（若他填在表里而不是正文行）
    if table_rows:
        for t in doc.tables:
            if len(t.rows) == table_rows and len(t.columns) == 4:
                for r in t.rows[1:]:
                    no = r.cells[0].text.strip()
                    if no in ids:
                        v = r.cells[3].text.strip()
                        if v and not sections.get(no, {}).get("choice"):
                            sections.setdefault(no, {"context": "", "recommend": "", "options": [],
                                                     "checked": [], "unchecked": [],
                                                     "choice": "", "notes": []})
                            sections[no]["choice"] = v

    return {"sections": sections, "global_notes": global_notes}

