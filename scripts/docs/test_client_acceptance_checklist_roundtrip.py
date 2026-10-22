# -*- coding: utf-8 -*-
"""《客户端验收清单》往返测试。

做四件事：
  1) 生成文档并做结构验收（zip 完整 / XML 可解析 / PAGE 域在位）；
  2) 注入勾选（通过 / 不通过 / 跳过）+ 一行问题记录 + 一条总批注；
  3) 回读并逐项断言读回来的结果与注入的**完全一致**
     —— 特别是「三个并列框」必须**逐框**判定，不能整行算勾上；
  4) **反向自测**：故意造一份不一样的版本，回读必须报出差异，
     否则这个读取器只是个「永远输出 OK」的假闸门。
"""
from __future__ import annotations

import os
import shutil
import sys
import tempfile

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(REPO, "scripts", "docs"))

from docx import Document  # noqa: E402

import build_client_acceptance_checklist as builder  # noqa: E402
import read_client_acceptance_checklist as reader  # noqa: E402
from _decision_form_common import split_boxes, verify, structure_ok  # noqa: E402

BOX_CHARS = "□☑√✓✔■▣☒"


def _set_result(doc, no, mark):
    """把某编号行的结果列改成：只有 mark 那个框被勾上。"""
    for t in doc.tables:
        head = [(c.text or "").strip() for c in t.rows[0].cells]
        if head[:2] != ["编号", "验收项"]:
            continue
        for r in t.rows[1:]:
            texts = [(c.text or "").strip() for c in r.cells]
            if len(set(texts)) == 1 or texts[0] != no:
                continue
            cell = r.cells[4]
            p = cell.paragraphs[0]
            for run in list(p.runs):
                run._element.getparent().remove(run._element)
            labels = ["通过", "不通过", "跳过"]
            for lab in labels:
                ch = "☑" if lab == mark else "□"
                p.add_run("%s %s　" % (ch, lab))
            return True
    return False


def _set_problem(doc, no, item, desc, severity, shot="shot-01.png"):
    for t in doc.tables:
        head = [(c.text or "").strip() for c in t.rows[0].cells]
        if head[:2] != ["编号", "对应验收项"]:
            continue
        for r in t.rows[1:]:
            if (r.cells[0].text or "").strip():
                continue
            r.cells[0].text = no
            r.cells[1].text = item
            r.cells[2].text = desc
            r.cells[3].text = "☑ %s　□ 阻塞　□ 小" % severity
            r.cells[4].text = shot
            return True
    return False


def _append_note(doc, text):
    doc.add_paragraph(text)


def main():
    tmpdir = tempfile.mkdtemp(prefix="acc_roundtrip_")
    docx_path = os.path.join(tmpdir, "checklist.docx")
    original_out = builder.OUT
    builder.OUT = docx_path
    try:
        # ── 1) 生成 + 结构验收 ────────────────────────────
        builder.build()
        info = verify(docx_path)
        print("[1] structure:", "PASS" if structure_ok(info) else "FAIL", info)
        if not structure_ok(info):
            return 1
        assert info["tables"] >= 2, "至少要有主表 + 问题记录表两张"
        assert info["footer_page_field"], "页脚 PAGE 域缺失"

        # ── 2) 注入 ───────────────────────────────────────
        doc = Document(docx_path)
        inject = [("A1", "通过"), ("D5", "不通过"), ("F5", "跳过")]
        for no, mark in inject:
            assert _set_result(doc, no, mark), "注入失败: %s" % no
        assert _set_problem(doc, "D5", "点同步到云端", "报 云端不可达 [WinError 10053]", "阻塞")
        _append_note(doc, "总规则：先别管生图，把发布链路跑通")
        doc.save(docx_path)
        print("[2] injected:", inject)

        # ── 3) 回读断言 ───────────────────────────────────
        r = reader.parse(docx_path)
        s = r["summary"]
        print("[3] readback:", s)
        by_no = {i["no"]: i for i in r["items"]}
        assert len(by_no) == s["total"], "条目数对不上"
        for no, mark in inject:
            assert by_no[no]["result"] == mark, \
                "%s 结果读成 %r，期望 %r" % (no, by_no[no]["result"], mark)
        assert s["passed"] == 1 and s["failed"] == 1 and s["skipped"] == 1, s
        assert s["failed_ids"] == ["D5"], s["failed_ids"]
        assert len(r["problems"]) == 1, r["problems"]
        pr = r["problems"][0]
        assert pr["no"] == "D5" and pr["severity"] == "阻塞", pr
        assert "WinError 10053" in pr["desc"], pr["desc"]
        assert pr["shot"] == "shot-01.png", pr["shot"]
        assert any("总规则" in n for n in r["global_notes"]), r["global_notes"]
        print("[3] PASS —— 逐框判定 / 问题记录 / 总批注 都读对了")

        # 并列框必须逐框判定：同一行三个框里只勾一个
        raw = by_no["D5"]["result_raw"]
        boxes = split_boxes(raw)
        checked = [lab for _p, lab, ck in boxes if ck]
        assert len(boxes) == 3, boxes
        assert checked == ["不通过"], boxes
        print("[3] 并列框逐框判定 OK ->", boxes)

        # ── 4) 反向自测：故意改内容，必须报出差异 ──────────
        doc2 = Document(docx_path)
        assert _set_result(doc2, "D5", "通过"), "反向注入失败"
        doc2.save(docx_path)
        r2 = reader.parse(docx_path)
        s2 = r2["summary"]
        print("[4] mutated readback:", s2)
        assert s2["failed"] == 0, "改成通过后 failed 应为 0，实际 %s" % s2
        assert "D5" not in s2["failed_ids"], s2["failed_ids"]
        assert s2["passed"] == 2, s2
        print("[4] PASS —— 读取器能报出差异，不是假闸门")
        return 0
    finally:
        builder.OUT = original_out
        shutil.rmtree(tmpdir, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main())
