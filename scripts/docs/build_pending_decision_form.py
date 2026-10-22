"""生成《Prism-Fusion 开源边界 · 第二轮待决表》Word 文档。

第一张拍板表（docs/oss-boundary-decision-form.docx）的 10 项总体决策 + 7 项运维待办
已写回仓库并落代码；本表只收剩下的 4 处尾巴：
  R1 D2 绑机器最终口径（第一张表里两个互斥项同时勾了）
  R2 D6 云端 prompt 移出开源树的实现方式
  R3 §1 55 项 + §2 54 项兜底结果是否要改
  R4 干净 Windows VM 验证 + 真实生产联调凭证

用法：
  D:\\dev\\python\\python.exe scripts\\docs\\build_pending_decision_form.py

依赖：python-docx（本机装在 D 盘 dev python，不写 C 盘、不动项目 venv）。
"""
import io
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

FONT_BODY = "等线"
FONT_HEAD = "微软雅黑"
C_DARK = RGBColor(0x1F, 0x38, 0x64)
C_MID = RGBColor(0x2E, 0x53, 0x95)
C_GRAY = RGBColor(0x59, 0x59, 0x59)
C_HDRBG = "DCE6F1"
C_LINE = "BFCEDD"

BOX = "□"   # □ 勾选框


# ----------------------------------------------------------------- 工具函数
def set_run_font(run, name=FONT_BODY, size=10.5, bold=False, color=None, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:eastAsia"), name)
    rf.set(qn("w:ascii"), name)
    rf.set(qn("w:hAnsi"), name)


def para(doc, text="", size=10.5, bold=False, color=None, align=None,
         space_before=0, space_after=4, indent=0, name=FONT_BODY, line=1.4, style=None):
    p = doc.add_paragraph(style=style)
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line
    if indent:
        pf.left_indent = Cm(indent)
    if align is not None:
        p.alignment = align
    if text:
        set_run_font(p.add_run(text), name=name, size=size, bold=bold, color=color)
    return p


def heading(doc, text, level=1):
    cfg = {1: (16, 14, 6, C_DARK, FONT_HEAD),
           2: (13, 12, 5, C_MID, FONT_HEAD),
           3: (11.5, 10, 4, RGBColor(0x33, 0x33, 0x33), FONT_HEAD)}
    size, sb, sa, color, font = cfg[level]
    p = para(doc, text, size=size, bold=True, color=color,
             space_before=sb, space_after=sa, name=font, line=1.2)
    return p


def shade(cell, hexcolor):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcpr.append(shd)


def cell_text(cell, text, size=9, bold=False, color=None, align=None, name=FONT_BODY):
    cell.text = ""
    p = cell.paragraphs[0]
    pf = p.paragraph_format
    pf.space_before = Pt(1.5)
    pf.space_after = Pt(1.5)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15
    if align is not None:
        p.alignment = align
    set_run_font(p.add_run(text), name=name, size=size, bold=bold, color=color)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    trpr.append(h)


def no_split(row):
    trpr = row._tr.get_or_add_trPr()
    c = OxmlElement("w:cantSplit")
    trpr.append(c)


def table_borders(table, color=C_LINE, sz=4):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:%s" % edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        borders.append(e)
    tbl_pr.append(borders)


def set_widths(table, widths):
    table.autofit = False
    table.allow_autofit = False
    for row in table.rows:
        for idx, w in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = Cm(w)


def add_page_footer(section, left_text):
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    set_run_font(p.add_run(left_text + "　·　第 "), size=8, color=C_GRAY)
    # PAGE 域：必须带一个缓存 run，否则部分 Word 版本会渲染成空白
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    r2 = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "16")
    rpr.append(sz)
    r2.append(rpr)
    t2 = OxmlElement("w:t")
    t2.text = "1"
    r2.append(t2)
    fld.append(r2)
    p._p.append(fld)
    set_run_font(p.add_run(" 页"), size=8, color=C_GRAY)


def blank_lines(doc, n=3, label=None):
    if label:
        para(doc, label, size=9.5, bold=True, color=C_MID, space_before=6, space_after=3)
    for _ in range(n):
        p = para(doc, "", size=10.5, space_after=6)
        ppr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "C9C9C9")
        pbdr.append(bottom)
        ppr.append(pbdr)



# ----------------------------------------------------------------- 本轮内容
OUT2 = os.path.join(REPO, "docs", "oss-boundary-round2-decision-form.docx")

DECISIONS2 = [
    dict(
        no="R1", title="D2　注册码要不要绑机器（第一张表里两个互斥项同时勾了）",
        ctx="你在第一张表里同时勾了「不绑（一码通用）」和「绑，但允许一次自助解绑」—— 这两条互斥，必须二选一。"
            "本轮我**先按「默认不绑；若启用绑定则允许自助解绑」落地**（把绑定开关留成一处常量，改起来只改一行）。",
        opts=[
            "A. 默认不绑机器（一码通用）—— 好卖、易分享，代价是容易被整包转卖",
            "B. 绑机器，但允许用户自助解绑（每月 1 次）—— 防共享与易转手的折中",
            "C. 绑机器，解绑走人工 / 客服 —— 防共享最强，售后成本最高",
        ],
        rec="建议先按 A 落地。等真卖出去了、看盗版和转卖情况，再决定要不要收紧到 B 或 C（代码里开关已经留好，收紧不用重写）。",
        mine="我的选择：",
    ),
    dict(
        no="R2", title="D6　cloud_service/llm_proxy/prompts.py 怎么「移出开源树」",
        ctx="这个文件 18KB，装的是云端提示词模板。你定了「移出开源树、只在 Pro 包提供」。"
            "但它现在**就在开源的 cloud_service/ 目录里**，还被 llm_proxy/router.py 与 db.py import —— "
            "直接删掉的话，公开仓 clone 下来 import 就炸。所以要你定**移的方式**。",
        opts=[
            "A. 公开仓只留 stub（接口在、真实模板不在，调用时明确提示「此功能需 Pro 版」），Pro 打包时注入真实文件",
            "B. 发布 GitHub 前用脚本剥离，仓库里保留完整文件",
            "C. 接受现状，继续开源（放弃「prompt 不可抄」这个差异化定位）",
        ],
        rec="建议 A。只有 A 能同时做到「能公开分发」和「公开仓 clone 下来不报错」；"
            "B 的保护是名义上的 —— GitHub 历史里照样能翻出模板。C 等于自己放弃定位。",
        mine="我的选择：",
    ),
    dict(
        no="R3", title="§1 55 项 + §2 54 项（共 109 项）的兜底结果要不要改",
        ctx="这 109 项你在第一张表里**没有逐项勾**（快速批注区也没写），我按每行的「当前建议」列兜底了 —— "
            "所以现在的结果是**推断，不是你逐项确认过的**。当前分布："
            "§1 开源 41 / 注册码 14 / 不做 0；§2 做入口且开源 23 / 做入口但注册码 16 / 不做入口 15。",
        opts=[
            "A. 就按兜底执行",
            "B. 我要改几项，在下面写「序号 + 新选择」（例：§1-14 代理管理 → 开源）",
            "C. 全部放开成开源（§1 那 14 项注册码全变开源）",
        ],
        rec="建议先按 A。等注册码模块（B1）做出来、界面上真能看见「哪些被锁」了，再按实际观感调 —— "
            "那时改就是看图说话，比现在拍脑袋准。",
        mine="我的选择（要改就写序号 + 新选择）：",
    ),
    dict(
        no="R4", title="干净 VM 验证 与 真实生产联调凭证（运维第 2、3 项，你没勾）",
        ctx="这两项没勾，意味着两件关键事目前**没有被验证过**："
            "①装到一台干净 Windows 机器上能不能从头跑通（现在两个安装器都早于 HEAD，用户拿到的是旧侧边栏）；"
            "②支付 / 云端 AI / TikHub 这条链是不是真的能用（现在只有 mock 跑通）。"
            "这部分我**没有你的环境就做不了**，需要你给路。",
        opts=[
            "VM 验证 → 推进（我需要一台干净 Windows 机器或虚拟机的操作权）",
            "VM 验证 → 暂缓",
            "生产凭证 → 我提供（下面写怎么给）",
            "生产凭证 → 改用沙箱 / 模拟通道先跑通",
            "生产凭证 → 暂缓",
        ],
        rec="凭证**不要贴在对话里**。最省事的做法：你在 D 盘建个目录，把 .env 文件放进去，"
            "我按路径读（不打印内容、不写进仓库、不进 git）。支付证书私钥同理，路径给我就行。",
        mine="凭证怎么给（写路径或方式）：",
    ),
]

doc2 = Document()
normal2 = doc2.styles["Normal"]
normal2.font.name = FONT_BODY
normal2.font.size = Pt(10.5)
normal2.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)

sec2 = doc2.sections[0]
sec2.page_width = Cm(21.0)
sec2.page_height = Cm(29.7)
sec2.top_margin = Cm(2.0)
sec2.bottom_margin = Cm(1.8)
sec2.left_margin = Cm(1.8)
sec2.right_margin = Cm(1.8)
add_page_footer(sec2, "Prism-Fusion 开源边界 · 第二轮待决表　v2.0　2026-10-22")

# ---------------- 封面
para(doc2, "", space_after=26)
para(doc2, "Prism-Fusion（老马AI工作台）", size=12.5, color=C_GRAY,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
para(doc2, "开源边界 · 第二轮待决表", size=25, bold=True, color=C_DARK,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8, name=FONT_HEAD)
para(doc2, "第一张表的 4 处尾巴 —— 勾完我就能开工", size=12, color=C_MID,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=22)
para(doc2, "版本 v2.0　|　日期 2026-10-22　|　对应代码 0753cc8", size=10,
     color=C_GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3)
para(doc2, "第一张表的 10 项总体决策 + 7 项运维待办已全部写回仓库（docs/feature-inventory-and-oss-boundary.md §7）",
     size=9.5, color=C_GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=26)

t = doc2.add_table(rows=0, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
table_borders(t)
info2 = [
    ("这份表解决什么", "第一张表已落地，只剩 4 处尾巴：D2 勾了互斥项、D6 移出开源树的实现方式、109 项兜底结果是否要改、干净 VM 与生产凭证要不要推进。"),
    ("一共几处要你拍板", "4 项（R1–R4）。每一项我都写了「我已核实的事实」「选项」「我的建议」，你只需要勾。"),
    ("填完怎么给我", "① 把这个 .docx 存回发给我；② 或者直接念勾选结果（「R1 选 A、R2 选 A、R3 选 A、R4 VM 暂缓、凭证走 D 盘目录」）。两种都行。"),
    ("填不完怎么办", "不想逐项勾，就在最下面的「总快速批注」区写一句总规则，例如「R1/R2 都按建议来」。"),
]
for k, v in info2:
    row = t.add_row()
    shade(row.cells[0], C_HDRBG)
    cell_text(row.cells[0], k, size=10, bold=True, color=C_DARK)
    cell_text(row.cells[1], v, size=10)
set_widths(t, [4.0, 12.6])

doc2.add_page_break()

# ---------------- §0 怎么填
heading(doc2, "0　怎么填这份表", 1)
how2 = doc2.add_table(rows=0, cols=2)
table_borders(how2)
for k, v in [
    ("选项行", "一行几个候选，前面带 □ 框。在框里打「√」，或者把 □ 改成 ☑（Word 里输入 2611 字符），也可以直接在框后写「选 A」。"),
    ("「我的选择」行", "每个小节最后那行「→ 我的选择：____」，直接写 A/B/C 或一句话。"),
    ("快速批注", "小节下面的横线区，或文末的总批注区。写一句总规则就够，我按它兜底。"),
]:
    row = how2.add_row()
    shade(row.cells[0], C_HDRBG)
    cell_text(row.cells[0], k, size=10, bold=True, color=C_DARK)
    cell_text(row.cells[1], v, size=10)
set_widths(how2, [3.4, 13.2])

# ---------------- §1 速览表
heading(doc2, "1　待决 4 项速览（一屏看完）", 1)
para(doc2, "「当前建议」是我基于已核实事实给的推荐，不是结论；「我的选择」列留空给你填。",
     size=10, color=C_GRAY, space_after=6)

cols = ["项", "要你定什么", "当前建议", "我的选择"]
w = [1.4, 6.6, 5.6, 3.0]
ts = doc2.add_table(rows=1, cols=4)
ts.style = "Table Grid"
table_borders(ts)
for i, c in enumerate(cols):
    cell = ts.rows[0].cells[i]
    shade(cell, C_HDRBG)
    cell_text(cell, c, size=9.5, bold=True, color=C_DARK,
              align=WD_ALIGN_PARAGRAPH.CENTER if i == 0 else None)
repeat_header(ts.rows[0])

SHORT2 = {
    "R1": "D2 绑机器最终口径",
    "R2": "D6 云端 prompt 怎么移出开源树",
    "R3": "109 项兜底结果要不要改",
    "R4": "干净 VM 验证 + 生产联调凭证",
}
for d in DECISIONS2:
    row = ts.add_row()
    no_split(row)
    cell_text(row.cells[0], d["no"], size=9.5, bold=True, color=C_DARK,
              align=WD_ALIGN_PARAGRAPH.CENTER)
    cell_text(row.cells[1], SHORT2[d["no"]], size=9)
    cell_text(row.cells[2], d["rec"], size=8.5)
    cell_text(row.cells[3], "", size=9)
set_widths(ts, w)

para(doc2, "", space_after=6)
para(doc2, "⚠️ R1 与 R2 会改代码（前者改注册码模块的绑定常量，后者改开源树的发布结构），"
           "R3 会改 109 行表格，R4 决定我能不能跑真机验证。", size=9.5, color=C_MID, space_after=8)

doc2.add_page_break()

# ---------------- §2 逐项详述
heading(doc2, "2　逐项详述", 1)
for d in DECISIONS2:
    heading(doc2, "%s　%s" % (d["no"], d["title"]), 3)
    if d["ctx"]:
        para(doc2, d["ctx"].replace("**", ""), size=9.5, color=C_GRAY, space_after=4)
    for o in d["opts"]:
        para(doc2, "%s %s" % (BOX, o), size=10, indent=0.4, space_after=3)
    if d["rec"]:
        para(doc2, "💡 " + d["rec"], size=9.5, color=C_MID, indent=0.4, space_after=3)
    para(doc2, "→ " + d["mine"] + "_______________________________________________",
         size=10, bold=True, space_before=2, space_after=6)
    blank_lines(doc2, 2, "小节批注（可留空）：")

doc2.add_page_break()

# ---------------- §3 总批注
heading(doc2, "3　总快速批注", 1)
para(doc2, "任何一项不想细勾，就在这里写一句总规则，我按它兜底（并会在回执里标明哪些是兜底的）。",
     size=10, color=C_GRAY, space_after=6)
blank_lines(doc2, 8)

heading(doc2, "4　填完之后我会做什么", 1)
for txt in [
    "① 读出你的选择 → 写回 docs/feature-inventory-and-oss-boundary.md §7.2 与 feature gate 矩阵；",
    "② R1 / R2 落到代码：注册码模块的绑定常量、cloud_service/llm_proxy/prompts.py 的发布结构；",
    "③ R3 若有改动 → 逐项改 109 行并重生成矩阵；",
    "④ R4 若推进 → 出干净 VM 验证清单 / 凭证读取方案（凭证不打印、不入库、不进 git）；",
    "⑤ 全部改动跑回归（pytest / tsc / vitest）后提交并推私有仓库 + 备份。",
]:
    para(doc2, txt, size=10, indent=0.3, space_after=4)

doc2.save(OUT2)
print("已生成:", OUT2, os.path.getsize(OUT2), "bytes")
